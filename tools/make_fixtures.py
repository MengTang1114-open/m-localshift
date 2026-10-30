from pathlib import Path
import subprocess,json,zipfile,tarfile,gzip,plistlib,struct,zlib,shutil,uuid
from PIL import Image,ImageDraw
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from reportlab.pdfgen import canvas
root=Path(__file__).resolve().parent.parent
out=root/'work/functional/fixtures';out.mkdir(parents=True,exist_ok=True)
failures=[]
def run(args):
 p=subprocess.run([str(a) for a in args],capture_output=True,text=True,timeout=90)
 if p.returncode:raise RuntimeError(p.stderr[-1000:] or p.stdout[-1000:])
 return p.stdout
im=Image.new('RGB',(320,180),'#1e433c');d=ImageDraw.Draw(im);d.rectangle((20,20,100,150),fill='#ff8a32');d.text((130,65),'LocalShift 2026',fill='white');im.save(out/'sample.png');im.save(out/'sample.webp')
ff=root/'LocalShift.app/Contents/Resources/bin/ffmpeg'
for ext in ['jpg','bmp','tiff','gif']:
 im.save(out/('sample.'+ext))
for ext,codec in [('heic',[]),('avif',['-c:v','libsvtav1','-crf','35','-preset','12'])]:
 try:
  if ext=='heic':continue
  run([ff,'-nostdin','-v','error','-i',out/'sample.png',*codec,'-frames:v','1','-y',out/('sample.'+ext)])
 except Exception as e:failures.append({'fixture':ext,'error':str(e)})
# System sips supplies HEIC, independent of the conversion implementation.
try:run(['/usr/bin/sips','-s','format','heic',out/'sample.png','--out',out/'sample.heic'])
except Exception as e:failures.append({'fixture':'heic','error':str(e)})
(out/'sample.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="320" height="180"><rect width="320" height="180" fill="#1e433c"/><circle cx="80" cy="90" r="45" fill="#ff8a32"/><text x="145" y="90" fill="white">LocalShift</text></svg>')
run([ff,'-nostdin','-v','error','-f','lavfi','-i','testsrc2=size=320x180:rate=12','-f','lavfi','-i','sine=frequency=440:sample_rate=44100','-t','1.2','-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac','-y',out/'sample.mp4'])
for ext,opts in [('mov',['-c:v','libx264','-c:a','aac']),('mkv',['-c:v','libx264','-c:a','aac']),('webm',['-c:v','libvpx-vp9','-c:a','libopus']),('avi',['-c:v','mpeg4','-c:a','mp3']),('wmv',['-c:v','wmv2','-c:a','wmav2'])]:
 run([ff,'-nostdin','-v','error','-i',out/'sample.mp4',*opts,'-y',out/('sample.'+ext)])
run([ff,'-nostdin','-v','error','-f','lavfi','-i','sine=frequency=440:sample_rate=44100','-t','1.2','-y',out/'sample.wav'])
for ext,codec in [('mp3','libmp3lame'),('m4a','aac'),('flac','flac'),('ogg','vorbis'),('opus','libopus'),('aiff','pcm_s16be'),('wma','wmav2')]:
 extra=['-strict','experimental','-ac','2'] if ext=='ogg' else []
 run([ff,'-nostdin','-v','error','-i',out/'sample.wav','-c:a',codec,*extra,'-y',out/('sample.'+ext)])
text='LocalShift local file conversion\n\nHello world. 本地文件转换。\nSecond paragraph with numbers 123.\n'
(out/'sample.txt').write_text(text);(out/'sample.md').write_text('# LocalShift\n\n'+text+'\n- First item\n- Second item\n')
doc=Document();doc.add_heading('LocalShift',0);doc.add_paragraph('Local conversion with editable text. 本地文件转换。');table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Item';table.cell(0,1).text='Value';table.cell(1,0).text='Sample';table.cell(1,1).text='123';doc.save(out/'sample.docx')
pres=Presentation();slide=pres.slides.add_slide(pres.slide_layouts[1]);slide.shapes.title.text='LocalShift';slide.placeholders[1].text='Editable slide text 123';pres.save(out/'sample.pptx')
wb=Workbook();ws=wb.active;ws.title='Local';ws.append(['Name','Value']);ws.append(['Sample',123]);ws.append(['中文',456]);wb.save(out/'sample.xlsx')
(out/'sample.csv').write_text('Name,Value\n"中文,引用",123\nSample,456\n');(out/'sample.tsv').write_text('Name\tValue\n中文\t123\nSample\t456\n')
c=canvas.Canvas(str(out/'sample.pdf'),pagesize=(320,180));c.drawString(25,140,'LocalShift editable PDF text 123');c.drawString(25,110,'Name                 Value');c.drawString(25,85,'Sample                123');c.showPage();c.drawString(25,140,'Second page');c.showPage();c.save()
for ext in ['rtf','rtfd']:
 try:run(['/usr/bin/textutil','-convert',ext,'-output',out/('sample.'+ext),out/'sample.txt'])
 except Exception as e:failures.append({'fixture':ext,'error':str(e)})
(out/'sample.textClipping').write_bytes(plistlib.dumps({'UTI-Data':{'public.utf8-plain-text':text.encode()}},fmt=plistlib.FMT_BINARY))
with zipfile.ZipFile(out/'sample.epub','w') as z:
 z.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
 z.writestr('META-INF/container.xml','<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
 z.writestr('OEBPS/content.opf','<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">localshift-test</dc:identifier><dc:title>LocalShift</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-10-27T00:00:00Z</meta></metadata><manifest><item id="page" href="page.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="page"/></spine></package>')
 z.writestr('OEBPS/page.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><head><title>LocalShift</title></head><body><h1>LocalShift</h1><p>Editable ebook text 123</p></body></html>')
 z.writestr('OEBPS/nav.xhtml','<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="page.xhtml">LocalShift</a></li></ol></nav></body></html>')
(out/'sample.srt').write_text('1\n00:00:00,100 --> 00:00:00,800\nLocalShift 本地转换\n\n');(out/'sample.vtt').write_text('WEBVTT\n\n00:00:00.100 --> 00:00:00.800\nLocalShift 本地转换\n\n')
folder=out/'sample-folder';folder.mkdir(exist_ok=True);(folder/'file.txt').write_text(text)
with zipfile.ZipFile(out/'sample.zip','w') as z:z.write(folder/'file.txt','sample-folder/file.txt')
with tarfile.open(out/'sample.tar','w') as t:t.add(folder,arcname='sample-folder')
with tarfile.open(out/'sample.gz','w:gz') as t:t.add(folder,arcname='sample-folder')
# A tiny valid RAR4 stored fixture; this checks the input container, not compressed-RAR codec breadth.
def block(payload):return struct.pack('<H',zlib.crc32(payload)&0xffff)+payload
name=b'file.txt';body=b'LocalShift RAR fixture\n'
main=block(struct.pack('<BHHHI',0x73,0,13,0,0))
fileheader=block(struct.pack('<BHHIIBIIBBHI',0x74,0x8000,32+len(name),len(body),len(body),3,zlib.crc32(body)&0xffffffff,0,20,0x30,len(name),0o100644)+name)
end=block(struct.pack('<BHH',0x7b,0x4000,7))
(out/'sample.rar').write_bytes(b'Rar!\x1a\x07\x00'+main+fileheader+body+end)
# Each LibreOffice call owns an isolated profile and only our fixture files.
lo=root/'LocalShift.app/Contents/Resources/OfficeRuntime/LibreOffice.app/Contents/MacOS/soffice'
for src,fmt in [('sample.xlsx','ods'),('sample.xlsx','xls'),('sample.pptx','ppt')]:
 try:run([lo,'-env:UserInstallation='+(out/('profile-'+fmt)).as_uri(),'--headless','--convert-to',fmt,'--outdir',out,out/src])
 except Exception as e:failures.append({'fixture':fmt,'error':str(e)})
# Use an operation-owned LaunchServices import to supply Pages sandbox access.
if not (out/'sample.pages').exists():
 try:
  unique='LocalShift-Fixture-'+str(uuid.uuid4())
  copy=out/(unique+'.docx');shutil.copy2(out/'sample.docx',copy)
  run(['/usr/bin/open','-g','-a','Pages',copy])
  run(['/usr/bin/osascript',root/'Resources/OfficePages.applescript',copy,out/'sample.pages','pages',unique])
  copy.unlink()
 except Exception as e:failures.append({'fixture':'pages','error':str(e)})
matrix=json.loads((root/'docs/conversion-matrix.json').read_text())['matrix'];jobs=[]
missing=[]
for src,targets in matrix.items():
 path=out/('sample.'+('textClipping' if src=='textclipping' else src))
 if not path.exists():missing.append(src);continue
 for ext in targets:jobs.append({'name':src+' → '+ext,'inputs':[str(path)],'option':'format_'+ext})
(root/'work/functional/matrix-jobs.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2))
# inspect_final_outputs.py addresses these results by index. Do not reorder.
photo=Image.new('RGB',(1800,1200),'#24564c')
pen=ImageDraw.Draw(photo)
for band in range(120):
    pen.rectangle((band*14,(band*37)%1100,band*14+90,(band*37)%1100+140),fill=((band*17)%255,(band*29)%255,(band*13)%255))
photo.save(out/'photo.jpg',quality=95)
if (out/'photo.jpg').stat().st_size<=40000:
    raise SystemExit('photo.jpg must be larger than the 40000-byte compression target')
shutil.copy2(out/'sample.pdf',out/'second.pdf')
def tool_job(name,option,filenames,**extra):
    item={'name':name,'option':option,'inputs':[str(out/filename) for filename in filenames]}
    item.update(extra)
    return item
tool_jobs=[
    tool_job('Image compress target','compress',['photo.jpg'],compressionMode='target',targetBytes=40000),
    tool_job('mp4 compress','compress',['sample.mp4']),
    tool_job('wav compress','compress',['sample.wav']),
    tool_job('pdf compress','compress',['sample.pdf']),
    tool_job('jpg edit metadata','removeMetadata',['sample.jpg']),
    tool_job('mp4 edit metadata','removeMetadata',['sample.mp4']),
    tool_job('mp3 edit metadata','removeMetadata',['sample.mp3']),
    tool_job('pdf edit metadata','removeMetadata',['sample.pdf']),
    tool_job('editImage','editImage',['sample.png'],exactSize=True),
    tool_job('annotateImage','annotateImage',['sample.png'],exactSize=True),
    tool_job('frameImage','frameImage',['sample.png'],exactSize=True),
    tool_job('cropImage','cropImage',['sample.png'],exactSize=True),
    tool_job('redactImage','redactImage',['sample.png'],exactSize=True),
    tool_job('normalizeAudio','normalizeAudio',['sample.wav']),
    tool_job('muteVideo','muteVideo',['sample.mp4']),
    tool_job('trimVideo','trimVideo',['sample.mp4']),
    tool_job('splitPDF selection duplicates','splitPDF',['sample.pdf']),
    tool_job('organizePDF duplicates','organizePDF',['sample.pdf']),
    tool_job('extract RAR','extractArchive',['sample.rar']),
    tool_job('Mixed images to PDF','createPDF',['sample.png','sample.jpg']),
    tool_job('mergePDF','mergePDF',['sample.pdf','second.pdf']),
    tool_job('cropVideo','cropVideo',['sample.mp4']),
    tool_job('changeVideoSpeed','changeVideoSpeed',['sample.mp4']),
    tool_job('videoSnapshots','videoSnapshots',['sample.mp4']),
    tool_job('splitVideo','splitVideo',['sample.mp4']),
    tool_job('redactVideo','redactVideo',['sample.mp4']),
    tool_job('joinVideos','joinVideos',['sample.mp4','sample.mov']),
    tool_job('audioToVideo','audioToVideo',['sample.wav']),
    tool_job('trimAudio','trimAudio',['sample.wav']),
    tool_job('audioChannels','audioChannels',['sample.wav']),
    tool_job('redactAudio','redactAudio',['sample.wav']),
    tool_job('createCollage','createCollage',['sample.png','sample.jpg']),
]
(root/'work/functional/tool-jobs.json').write_text(json.dumps(tool_jobs,ensure_ascii=False,indent=2))
report={'fixtures':len(matrix)-len(missing),'missing':missing,'jobs':len(jobs),'tool_jobs':len(tool_jobs),'generation_errors':failures,'notes':['RAR fixture uses stored entries','textClipping fixture generated as binary plist; real Finder clipping not yet verified']}
(root/'work/functional/fixture-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
