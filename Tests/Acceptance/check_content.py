from pathlib import Path
from decimal import Decimal
import csv,json,re,unicodedata,zipfile,subprocess,argparse
from html.parser import HTMLParser
from xml.etree import ElementTree as ET
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document
from PIL import Image
root=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser();parser.add_argument('--area',type=Path,default=root/'work/release-audit');parser.add_argument('--fixtures',type=Path,default=root/'work/release-audit/fixtures');args=parser.parse_args();area=args.area;fixtures=args.fixtures
cases=json.loads((area/'edge/results.json').read_text())['cases']; results=[]
norm=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',str(s)))
def output(i):return Path(cases[i]['outputs'][0])
def check(name,condition,detail=None):results.append({'name':name,'passed':bool(condition),'detail':detail})
def pdftext(i):return subprocess.check_output(['/opt/homebrew/bin/pdftotext','-layout',str(output(i)),'-'],text=True)
class VisibleHTML(HTMLParser):
    def __init__(self):super().__init__();self.parts=[]
    def handle_data(self,data):self.parts.append(data)
def visible_html(text):
    parser=VisibleHTML();parser.feed(text);return ''.join(parser.parts)
def present(text,phrases):return [p for p in phrases if norm(p) not in norm(text)]
base=json.loads((fixtures/'baseline.json').read_text());pdf=PdfReader(output(0));missing=present(pdftext(0),base['rich_phrases'])
check('Rich DOCX PDF retains Chinese Japanese Korean Latin and all page/table sentinels',not missing,missing)
check('Rich DOCX PDF explicit page break retained',len(pdf.pages)==2,len(pdf.pages))
check('Rich DOCX PDF embedded image retained',any(list(p.images) for p in pdf.pages))
for i in [1,2]:
 missing=present(output(i).read_text(),base['rich_phrases']);check(cases[i]['name']+' all text retained',not missing,missing)
html=next(output(3).glob('*.html'));missing=present(visible_html(html.read_text()),base['rich_phrases']);check('Rich DOCX HTML all text retained',not missing,missing)
check('Rich DOCX HTML image assets retained',bool(list(output(3).glob('*.png'))+list(output(3).glob('*.jpg'))))
with zipfile.ZipFile(output(4)) as z:
 text=''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xhtml'));missing=present(text,base['rich_phrases']);check('DOCX EPUB all text retained',not missing,missing)
check('DOCX PNG exports both pages',len(cases[5]['outputs'])==2)
slidepdf=PdfReader(output(6));missing=present(pdftext(6),base['slide_phrases']);check('PPTX PDF both slides and Chinese preserved',len(slidepdf.pages)==2 and not missing,missing)
check('PPTX PDF images preserved on both slides',all(list(p.images) for p in slidepdf.pages))
check('PPTX PNG exports both slides',len(cases[7]['outputs'])==2)
missing=present(pdftext(8),['中文数据','123.45','00123','133.45','尾页数据','789']);check('XLSX PDF both sheets calculated values and Chinese preserved',not missing,missing)
with zipfile.ZipFile(output(9)) as z:
 xml=z.read('content.xml').decode();check('XLSX ODS sheet names and formula retained',all(x in xml for x in ['销售明细','第二页','SUM','中文数据','00123']))
expected_sheets=[[['项目','金额'],['中文数据','123.45'],['保留前导零','00123'],['SUM','133.45']],[['尾页数据','789']]]
for i,delimiter in [(10,','),(11,'\t')]:
 actual_sheets=[]
 for name in cases[i]['outputs']:
  with Path(name).open(newline='') as f:actual_sheets.append(list(csv.reader(f,delimiter=delimiter)))
 check(cases[i]['name']+' all sheets retained as complete delimited files',actual_sheets==expected_sheets,actual_sheets)
rows=base['csv_rows']
T='{urn:oasis:names:tc:opendocument:xmlns:text:1.0}'
O='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
def ods_text(element):
 value=element.text or ''
 for child in element:
  if child.tag==T+'s':value+=' '*int(child.attrib.get(T+'c','1'))
  elif child.tag==T+'tab':value+='\t'
  elif child.tag==T+'line-break':value+='\n'
  else:value+=ods_text(child)
  value+=child.tail or ''
 return value
with zipfile.ZipFile(output(13)) as z:
 table=ET.fromstring(z.read('content.xml')).find('.//'+O+'table')
 got=[['\n'.join(ods_text(p) for p in cell.findall(T+'p')) for cell in row.findall(O+'table-cell')] for row in table.findall(O+'table-row')]
 check('CSV ODS exact cells and all decimal digits retained',got==rows,got)
missing=present(pdftext(15),[cell for row in rows for cell in row if cell])
check('CSV PDF all complete cell strings retained without clipping',not missing,missing)
wb=load_workbook(output(12));actual=list(wb.active.values)
check('CSV XLSX cells and quoted newlines retained',len(actual)==len(rows) and actual[1][0]==rows[1][0])
check('CSV XLSX leading zero ID remains text',actual[1][1]=='00123',str(actual[1][1]))
check('CSV XLSX arbitrary precision decimal remains exact',str(actual[2][1])==rows[2][1],str(actual[2][1]))
check('CSV XLSX long integer remains exact',str(actual[3][1])==rows[3][1],str(actual[3][1]))
check('CSV formula-like literal is never executed',actual[4][1]=='=1+2' and wb.active['B5'].data_type=='s',{'value':str(actual[4][1]),'type':wb.active['B5'].data_type})
check('CSV ordinary numeric amount remains numeric',isinstance(actual[5][1],(int,float)) and Decimal(str(actual[5][1]))==Decimal('123.45'))
with output(14).open(newline='') as f:got=list(csv.reader(f,delimiter='\t'))
check('CSV TSV roundtrip retains all exact cells',got==rows)
unicode_phrases=['中文标题','日本語の文字確認','한글 변환 확인','Résumé café naïve','# literal heading','a ** b','尾行不能丢失 END_SENTINEL']
for i in [16,17,18,19,20]:
 p=output(i)
 if p.suffix=='.pdf':text=pdftext(i)
 elif p.suffix=='.docx':text='\n'.join(x.text for x in Document(p).paragraphs)
 elif p.suffix=='.epub':
  with zipfile.ZipFile(p) as z:text=''.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xhtml'))
 else:text=p.read_text()
 missing=present(text,unicode_phrases);check(cases[i]['name']+' multilingual and literal code retained',not missing,missing)
report={'passed':sum(x['passed'] for x in results),'failed':sum(not x['passed'] for x in results),'cases':results,'scope':'Independent content oracles; checks content, page/sheet count, image presence, exact numeric values, formula safety. Does not imply arbitrary layout fidelity.'}
(area/'content-semantics.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:v for k,v in report.items() if k!='cases'},ensure_ascii=False));print(json.dumps([r for r in results if not r['passed']],ensure_ascii=False,indent=2));raise SystemExit(bool(report['failed']))
