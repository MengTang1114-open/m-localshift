from pathlib import Path
import json,stat,zipfile
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI
from openpyxl import Workbook
from PIL import Image,ImageDraw
root=Path(__file__).resolve().parents[2]
out=root/'work/release-audit/fixtures';out.mkdir(parents=True,exist_ok=True)
image=Image.new('RGB',(240,120),'#007aff');draw=ImageDraw.Draw(image);draw.rectangle((0,0,60,120),fill='#ff4e44');image.save(out/'embedded.png')
phrases=['中文内容不能丢失','日本語の文字確認','한글 변환 확인','Résumé café naïve','FIRST_PAGE_SENTINEL_123','TABLE_ROW_中文_456','LAST_PAGE_SENTINEL_789']
doc=Document();doc.add_heading('LocalShift 内容验收',0)
for text in phrases[:5]:doc.add_paragraph(text)
p=doc.add_paragraph();p.add_run('Bold text').bold=True;p.add_run(' and italic text').italic=True
for text in ['First list item','Second list item']:doc.add_paragraph(text,style='List Bullet')
t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='名称';t.cell(0,1).text='值';t.cell(1,0).text=phrases[5];t.cell(1,1).text='123.45'
doc.add_picture(str(out/'embedded.png'),width=Inches(2));doc.add_page_break();doc.add_heading('第二页',1);doc.add_paragraph(phrases[-1]);doc.save(out/'rich.docx')
pres=Presentation()
for title,body in [('中文幻灯片一','SLIDE_ONE_123 中文、日本語、한글'),('中文幻灯片二','SLIDE_TWO_456 尾页内容')]:
 slide=pres.slides.add_slide(pres.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body;slide.shapes.add_picture(str(out/'embedded.png'),PI(1),PI(4),width=PI(2))
pres.save(out/'slides.pptx')
wb=Workbook();ws=wb.active;ws.title='销售明细';ws.append(['项目','金额']);ws.append(['中文数据',123.45]);ws.append(['保留前导零','00123']);ws.append(['SUM','=SUM(B2,10)']);ws=wb.create_sheet('第二页');ws.append(['尾页数据',789]);wb.save(out/'sheets.xlsx')
# Values that must remain strings/numbers across table formats; trailing precision cannot disappear.
rows=[['name','value'],['中文,多行\n内容','00123'],['decimal','0.1234567890123456789'],['large','12345678901234567890'],['literal_formula','=1+2'],['amount','123.45']]
import csv
with (out/'precision.csv').open('w',newline='') as f:csv.writer(f).writerows(rows)
(out/'unicode.md').write_text('# 中文标题\n\n日本語の文字確認\n\n한글 변환 확인\n\nRésumé café naïve\n\n```text\n# literal heading\na ** b\n```\n\n尾行不能丢失 END_SENTINEL\n')
(out/'baseline.json').write_text(json.dumps({'rich_phrases':phrases,'slide_phrases':['SLIDE_ONE_123','中文幻灯片一','SLIDE_TWO_456','尾页内容'],'csv_rows':rows},ensure_ascii=False,indent=2))
jobs=[]
for file,targets in [('rich.docx',['pdf','txt','md','html','epub','png']),('slides.pptx',['pdf','png']),('sheets.xlsx',['pdf','ods','csv','tsv']),('precision.csv',['xlsx','ods','tsv','pdf']),('unicode.md',['docx','pdf','html','txt','epub'])]:
 for target in targets:jobs.append({'name':file+' → '+target,'inputs':[str(out/file)],'option':'format_'+target})
(root/'work/release-audit/edge-jobs.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2))
negative=root/'work/release-audit/negative-fixtures'
negative.mkdir(parents=True,exist_ok=True)
payload=b'not a valid document or image\x00\xff'
for name in ['corrupt.png','corrupt.pdf','corrupt.docx','corrupt.zip','corrupt.mp4']:
    (negative/name).write_bytes(payload)
Image.new('RGB',(20000,1),'#ff4e44').save(negative/'oversized.png')
(negative/'external.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="https://example.invalid/never-fetch"/></svg>')
with zipfile.ZipFile(negative/'traversal.zip','w') as archive:
    archive.writestr('../outside-marker.txt','never escape')
with zipfile.ZipFile(negative/'symlink.zip','w') as archive:
    info=zipfile.ZipInfo('escape')
    info.create_system=3
    info.external_attr=(stat.S_IFLNK|0o777)<<16
    archive.writestr(info,'../../target')
with zipfile.ZipFile(negative/'macro.docx','w') as archive:
    archive.writestr('word/document.xml','<document/>')
    archive.writestr('word/vbaProject.bin','macro')
with zipfile.ZipFile(negative/'external.docx','w') as archive:
    archive.writestr('word/document.xml','<document/>')
    archive.writestr('word/_rels/document.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://example.invalid/never-fetch" TargetMode="External"/></Relationships>')
print(f'{len(jobs)} edge content cases prepared')
print(f'{len(list(negative.iterdir()))} negative fixtures prepared')
